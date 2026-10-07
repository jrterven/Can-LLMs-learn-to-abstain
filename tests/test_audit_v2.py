"""Synthetic CPU tests: fixed sample, blinding, XLSX integrity and returned labels."""
import importlib.util
from pathlib import Path
import sys
import zipfile

from openpyxl import load_workbook
import pytest

SPEC = importlib.util.spec_from_file_location("audit_v2_tests_module", Path(__file__).parents[1] / "analysis/audit_v2.py")
a = importlib.util.module_from_spec(SPEC); sys.modules[SPEC.name] = a; SPEC.loader.exec_module(a)


def units_fixture():
    units = {}
    pairs = {
        "critic_only": [(False, True)], "original_only": [(True, False)],
        "both_directions": [(False, True), (True, False)],
        "agreement_emit": [(True, True)], "agreement_reject": [(False, False)],
    }
    for split in a.SPLITS:
        for pattern in pairs:
            for label in ("correct", "error", "abstain"):
                for i in range(20):
                    key = f"{split}/{pattern}/{label}/{i}"
                    units[key] = {"unit_key": key, "question_id": key, "question": "Question?", "response": "Answer",
                        "split": split, "aliases": ["Reference"], "truncated": False,
                        "automatic_outcome": label, "automatic_reason": "fixture", "rl_representative": None,
                        "occurrences": [{"source_file": "PRIVATE_SENTINEL", "seed": 17}],
                        "selector_pairs": [{"original_emits": x, "critic_emits": y} for x, y in pairs[pattern]]}
        for arm in a.ARMS:
            for label in ("correct", "error", "abstain"):
                for i in range(10):
                    key = f"{split}/{arm}/{label}/{i}"
                    units[key] = {"unit_key": key, "question_id": key, "question": "Question?", "response": "Answer",
                        "split": split, "aliases": ["Reference"], "truncated": False,
                        "automatic_outcome": label, "automatic_reason": "fixture",
                        "rl_representative": {"arm": arm, "seed": a.SEEDS[i%3]}, "selector_pairs": [],
                        "occurrences": [{"source_file": "PRIVATE_SENTINEL", "seed": 17}]}
    return units


@pytest.fixture
def exported(tmp_path):
    out = tmp_path / "audit"
    a.export_audit(out, loader=lambda: (units_fixture(), {}))
    return out


def filled_workbook(exported, tmp_path, name="returned.xlsx"):
    book = load_workbook(exported / "auditoria-v2-150-respuestas.xlsx")
    for row in book["Auditoria"].iter_rows(min_row=2):
        row[5].value = "error"
    p = book["Procedencia"]
    for i, value in enumerate(("Anotador-1", "2026-10-04", "no", "", "si", ""), 2):
        p.cell(i, 2).value = value
    path = tmp_path / name; book.save(path)
    return path


def test_selection_is_deterministic_unique_stratified_and_has_weights():
    units = units_fixture()
    chosen, info = a.sample_units(units)
    reverse, info2 = a.sample_units(dict(reversed(list(units.items()))))
    assert [u["unit_key"] for u in chosen] == [u["unit_key"] for u in reverse]
    assert info == info2 and len({u["unit_key"] for u in chosen}) == 150
    for split in a.SPLITS:
        assert info["sample_by_block"][f"{split}/shared_disagreement"] == 35
        assert info["sample_by_block"][f"{split}/shared_agreement"] == 15
        assert info["sample_by_block"][f"{split}/rl"] == 25
    for u in chosen:
        s = u["sampling"]
        assert s["nominal_inclusion_fraction"] == s["n_h"]/s["N_h"]
        assert s["design_weight"] == s["N_h"]/s["n_h"]
    for split in a.SPLITS:
        patterns = {u["sampling"]["stratum"][2] for u in chosen if u["split"] == split and u["selector_pairs"]}
        assert {"critic_only", "original_only", "both_directions"} <= patterns


def test_shortage_aborts_instead_of_retuning():
    units = {k:v for k,v in units_fixture().items() if not k.startswith("test_nq/agreement")}
    with pytest.raises(ValueError, match="fixed quota"):
        a.sample_units(units)


def test_shared_frame_takes_precedence_over_duplicate_rl_occurrences():
    units = units_fixture()
    u = next(v for v in units.values() if v["selector_pairs"])
    before = a.frame_and_stratum(u)
    u["rl_representative"] = {"arm": a.ARMS[0], "seed": 17}
    assert a.frame_and_stratum(u) == before


def test_duplicate_occurrences_and_truncation_keep_distinct_units():
    ref = {"id":"q", "question":"Capital?", "aliases":["Paris"]}; units = {}
    row = {**ref, "text":"Paris", "truncated":False, "outcome":"correct", "reason":"exact_match"}
    for seed in (17,29):
        a.add_occurrence(units,row,ref,{"split":"test_trivia","arm":a.ARMS[0],"seed":seed},True)
    changed = {**row,"truncated":True,"outcome":"error","reason":"truncated"}
    a.add_occurrence(units,changed,ref,{"split":"test_trivia","arm":a.ARMS[1],"seed":43},True)
    assert len(units)==2 and sorted(len(u["occurrences"]) for u in units.values())==[1,2]
    assert a.unit_key(row)!=a.unit_key(changed)
    copy = {}; occurrences = list(reversed(next(u for u in units.values() if not u["truncated"])["occurrences"]))
    for occ in occurrences:
        a.add_occurrence(copy,row,ref,occ,True)
    assert copy[a.unit_key(row)]["rl_representative"]==units[a.unit_key(row)]["rl_representative"]
    with pytest.raises(ValueError,match="regrade"):
        a.add_occurrence({}, {**changed,"outcome":"correct"},ref,occ,True)


def test_pattern_direction_ignores_correctness_and_filter_rejects_invalid():
    assert a.decision_pattern([{"critic_emits":True,"original_emits":False}])=="critic_only"
    row={"text":"IDK","truncated":False}
    assert not a.emits(row,100,.65)
    assert not a.emits({"text":"Paris","truncated":True},100,.65)
    assert a.emits({"text":".IDK","truncated":False},100,.65)
    assert not a.emits({"text":"Paris","truncated":False},0,.5)


def test_export_has_blank_labels_and_no_private_leakage(exported):
    book=load_workbook(exported/"auditoria-v2-150-respuestas.xlsx",data_only=False)
    a.verify_workbook_structure(book)
    assert book["Auditoria"].max_row==151
    assert all(c.value is None for row in book["Auditoria"].iter_rows(min_row=2,min_col=6) for c in row)
    assert all(r[1].value is None for r in book["Procedencia"].iter_rows(min_row=2))
    with zipfile.ZipFile(exported/"auditoria-v2-150-respuestas.xlsx") as z:
        text="\n".join(z.read(n).decode() for n in z.namelist() if n.endswith(".xml"))
    for secret in ("PRIVATE_SENTINEL", *a.ARMS, "critic_only", "automatic_outcome", "selector_pairs"):
        assert secret not in text
    assert len(book["Auditoria"].data_validations.dataValidation)==2
    with pytest.raises(ValueError,match="Missing or invalid semantic"):
        a.validate_submission(exported/"auditoria-v2-150-respuestas.xlsx",exported)


def test_formula_like_dataset_strings_remain_literal(tmp_path):
    record={k:"" for k in a.FIELDS}
    record.update(audit_id="x",question="=1+1",response="@formula",reference_aliases="+Reference",truncated="no")
    path=tmp_path/"literal.xlsx";a.write_workbook(path,[record])
    book=load_workbook(path,data_only=False)
    assert book["Auditoria"]["B2"].value=="=1+1" and book["Auditoria"]["B2"].data_type=="s"
    assert a.csv_literal("=1+1")=="'=1+1"
    with pytest.raises(FileExistsError):a.write_workbook(path,[record])


def test_export_never_overwrites(exported):
    before={p.name:a.file_digest(p) for p in exported.iterdir()}
    with pytest.raises(FileExistsError):
        a.export_audit(exported,loader=lambda:pytest.fail("should not load sources"))
    assert before=={p.name:a.file_digest(p) for p in exported.iterdir()}


@pytest.mark.parametrize("change", ["source", "unknown_id", "duplicate_id", "label", "missing", "formula", "hidden", "unclear", "provenance"])
def test_import_rejects_bad_integrity_or_incomplete_annotations(exported,tmp_path,change):
    path=filled_workbook(exported,tmp_path)
    book=load_workbook(path);sheet=book["Auditoria"]
    if change=="source":sheet["C2"]="different"
    elif change=="unknown_id":sheet["A2"]="unknown"
    elif change=="duplicate_id":sheet["A2"]=sheet["A3"].value
    elif change=="label":sheet["F2"]="maybe"
    elif change=="missing":sheet.delete_rows(151)
    elif change=="formula":sheet["I2"]="=2+2"
    elif change=="hidden":sheet.column_dimensions["B"].hidden=True
    elif change=="unclear":sheet["F2"]="unclear"
    else:book["Procedencia"]["B4"]=""
    book.save(path)
    with pytest.raises(ValueError):a.import_audit(path,exported)
    assert not (exported/"submissions").exists()


def test_valid_return_preserves_semantic_labels_and_optional_format(exported,tmp_path):
    path=filled_workbook(exported,tmp_path)
    book=load_workbook(path);book["Auditoria"]["F2"]="abstain";book["Auditoria"]["G2"]="abstention_variant"
    book["Auditoria"]["F3"]="unclear";book["Auditoria"]["I3"]="Contexto temporal insuficiente."
    book["Procedencia"]["B4"]="si";book["Procedencia"]["B5"]="Asistente como búsqueda; revisé las fuentes."
    book.save(path)
    before={p.name:a.file_digest(p) for p in exported.iterdir()}
    result=a.import_audit(path,exported)
    assert result["n"]==150 and result["unclear"]==1 and result["ai_assistance_declared"] is True
    assert result["personal_review_all_declared"] is True and result["frozen_endpoints_modified"] is False
    assert all(a.file_digest(exported/name)==sha for name,sha in before.items())
    with pytest.raises(FileExistsError):a.import_audit(path,exported)


def test_reordered_rows_are_allowed_but_export_manifest_is_immutable(exported,tmp_path):
    path=filled_workbook(exported,tmp_path);book=load_workbook(path);sheet=book["Auditoria"]
    first=[c.value for c in sheet[2]];last=[c.value for c in sheet[151]]
    for col,value in enumerate(first,1):sheet.cell(151,col).value=value
    for col,value in enumerate(last,1):sheet.cell(2,col).value=value
    book.save(path)
    assert len(a.validate_submission(path,exported)[0])==150
    with (exported/"blind.csv").open("a") as f:f.write("changed")
    with pytest.raises(ValueError,match="export changed"):a.validate_submission(path,exported)


def test_ai_assistance_requires_provenance_description(exported,tmp_path):
    path=filled_workbook(exported,tmp_path);book=load_workbook(path)
    book["Procedencia"]["B4"]="si";book.save(path)
    with pytest.raises(ValueError,match="Describe the AI"):a.validate_submission(path,exported)
