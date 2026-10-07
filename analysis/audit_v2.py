"""Post-hoc blind semantic audit of v2: export blank XLSX; validate returned labels.

CPU only. This never changes frozen predictions, graders, rewards or endpoints.
Private outputs include sampling strata and links to every eligible occurrence;
publish this code, not the audit data. Requires openpyxl==3.1.5 for XLSX support.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import date
import json
import math
from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "experiments/v2-20261002"
DEFAULT_OUT = ROOT / "artifacts/audit-v2-20261004"
sys.path.insert(0, str(ROOT / "src"))
from abstention.io import digest, file_digest, read_json, read_jsonl, utcnow, write_json
from abstention.scoring import grade

SEED = 20261004
SEEDS = (17, 29, 43)
TAUS = (.65, .85)
ARMS = ("a_g4_single_tau", "b_g8_single_tau", "c_g8_paired_tau", "d_g8_paired_fixed075")
SPLITS = {"test_trivia": 2000, "test_nq": 500}
LABELS = ("correct", "error", "abstain", "unclear")
FORMATS = ("ok", "incomplete", "empty", "abstention_variant", "other", "unclear")
FIELDS = ("audit_id", "question", "response", "reference_aliases", "truncated",
          "human_outcome", "format_note", "evidence", "notes")
PROVENANCE = ("Identificador del anotador", "Fecha de revisión (AAAA-MM-DD)",
              "¿Usaste IA como apoyo? (si/no)", "Herramienta y uso, si corresponde",
              "¿Revisaste personalmente cada etiqueta? (si/no)", "Limitaciones o procedimiento (opcional)")
POLICY = {
    "seed": SEED, "size": 150,
    "timing": "Post-hoc diagnostic audit designed after final aggregate v2 results and before new human labels; not a preregistered confirmatory endpoint.",
    "unit": "Exact (question_id, response_text, truncated); distinct truncation flags remain distinct units.",
    "shared_frame": "One frozen original forced candidate for each held-out question, shared by original calibrated filter and three calibrated critics.",
    "rl_frame": "Distinct units emitted in primary-threshold RL predictions; exclude ALL units in the shared-candidate frame, not only selected candidates.",
    "per_dataset": {"shared_disagreement": 35, "shared_agreement": 15, "rl": 25},
    "decision_pattern": "Across the six seed×primary-threshold comparisons: critic_only, original_only, both_directions, agreement_emit or agreement_reject. Directions are assigned without reference to correctness.",
    "shared_strata": "dataset × disagreement/agreement block × decision pattern × strict automatic label",
    "rl_strata": "dataset × hash-chosen representative RL arm × strict automatic label; representative fixed before strata using all eligible RL occurrences",
    "allocation": "Hash-order nonempty strata; round-robin allocate each fixed block quota, skipping exhausted strata. Select smallest independently salted unit hashes in each stratum.",
    "question_diversity": "No additional uniqueness-of-question filter; units are unique, questions may recur. This preserves simple stratum-conditional inclusion fractions.",
    "inclusion": "Retain N_h, n_h, nominal n_h/N_h and N_h/n_h for every selected unit. They refer to these explicitly restricted unit frames, not to all model outputs.",
    "shortage": "Abort on insufficient capacity in any of six dataset×block cells; do not retune using human labels.",
    "occurrences": "Track all direct forced/primary occurrences and all six selector-vs-original decision pairs for selected shared units. Diagnostic thresholds are outside this audit frame.",
    "blind": "Hide model, method, seed, threshold, automatic label, decision pattern, stratum and source IDs. Show question, exact response, fallible aliases, observed truncation flag and an opaque audit ID.",
    "labels": "Blank semantic labels correct/error/abstain/unclear; optional format note kept separate. No LLM annotation or inherited labels/provenance.",
    "limits": "Raw disagreement is not population accuracy, FNR or corrected utility. This audit does not replace frozen EM endpoints or establish a new method ranking. Any weighted sensitivity analysis requires explicit treatment of unresolved labels and repeated occurrences.",
}


def unit_key(row):
    return digest([row["id"], row["text"], row["truncated"]])


def sigmoid(z):
    if not math.isfinite(z):
        raise ValueError("Nonfinite confidence logit")
    return 1 / (1 + math.exp(-z)) if z >= 0 else math.exp(z) / (1 + math.exp(z))


def emits(row, logit, tau):
    return bool(row["text"].strip()) and not row["truncated"] and row["text"].strip() != "IDK" and sigmoid(logit) > tau


def decision_pattern(pairs):
    positive = any(p["critic_emits"] and not p["original_emits"] for p in pairs)
    negative = any(p["original_emits"] and not p["critic_emits"] for p in pairs)
    if positive and negative:
        return "both_directions"
    if positive:
        return "critic_only"
    if negative:
        return "original_only"
    return "agreement_emit" if any(p["original_emits"] for p in pairs) else "agreement_reject"


def add_occurrence(units, row, reference, occurrence, rl_eligible=False):
    if any(row[k] != reference[k] for k in ("id", "question")) or row.get("aliases", reference["aliases"]) != reference["aliases"]:
        raise ValueError("Prediction/reference mismatch")
    if not isinstance(row["text"], str) or type(row["truncated"]) is not bool:
        raise ValueError("Invalid response or truncation flag")
    automatic = grade(row["text"], reference["aliases"], row["truncated"])
    if (row["outcome"], row["reason"]) != (automatic.outcome, automatic.reason):
        raise ValueError("Frozen exact-match regrade mismatch")
    key = unit_key(row)
    if key not in units:
        units[key] = {"unit_key": key, "question_id": row["id"], "split": occurrence["split"],
                      "question": reference["question"], "response": row["text"], "aliases": reference["aliases"],
                      "truncated": row["truncated"], "automatic_outcome": automatic.outcome,
                      "automatic_reason": automatic.reason, "occurrences": [], "selector_pairs": [],
                      "rl_representative": None}
    unit = units[key]
    if unit["split"] != occurrence["split"]:
        raise ValueError("Question ID occurs in two datasets")
    unit["occurrences"].append(occurrence)
    if rl_eligible:
        old = unit["rl_representative"]
        if old is None or digest([SEED, "representative", occurrence]) < digest([SEED, "representative", old]):
            unit["rl_representative"] = occurrence
    return unit


def load_population():
    """Only verified completed v2 test outputs; no model loading or fitting."""
    artifacts = WORK / "artifacts"
    sources = {}
    def load(path, jsonl=False, expected=None):
        checksum = file_digest(path)
        if expected is not None and expected != checksum:
            raise ValueError(f"Changed source: {path}")
        sources[str(path.relative_to(ROOT))] = checksum
        return read_jsonl(path) if jsonl else read_json(path)
    variants = ["original", *(f"{arm}-s{s}" for arm in ARMS for s in SEEDS)]
    receipts = [artifacts / "predictions" / v / "complete.json" for v in variants]
    receipts += [artifacts / f"selector/test-evaluation/s{s}/complete.json" for s in SEEDS]
    if any(not p.is_file() for p in receipts):
        raise RuntimeError("All thirteen direct and three selector test completions are required")
    complete = load(artifacts / "analysis/complete.json")
    if complete["status"] != "artifacts_ready_for_review":
        raise ValueError("Final v2 analysis is not complete")
    verified = load(artifacts / "reviews/2026-10-04-completion/scientific.json")
    if verified["status"] != "passed":
        raise ValueError("Independent final scientific review has not passed")
    stats_path = artifacts / "analysis/statistics.json"
    stats_sha = file_digest(stats_path)
    if stats_sha != complete["outputs"]["statistics.json"] or stats_sha != verified["source_sha256"][str(stats_path.relative_to(ROOT))]:
        raise ValueError("Analysis differs from its independent review")
    sources[str(stats_path.relative_to(ROOT))] = stats_sha
    lock_path = artifacts / "evaluation-lock.json"
    lock = load(lock_path)
    if lock["status"] != "ready":
        raise ValueError("Evaluation lock is not ready")
    lock_sha = sources[str(lock_path.relative_to(ROOT))]
    data = {}
    for split, count in SPLITS.items():
        spec = lock["data"][split]
        rows = load(ROOT / spec["path"], True, spec["sha256"])
        if len(rows) != count or len({r["id"] for r in rows}) != count:
            raise ValueError("Unexpected test population")
        data[split] = rows
    units, shared = {}, {}
    for variant in variants:
        folder = artifacts / "predictions" / variant
        done = load(folder / "complete.json")
        if done["status"] != "complete" or done["variant"] != variant or done["evaluation_lock_sha256"] != lock_sha:
            raise ValueError("Incomplete or unrelated direct prediction receipt")
        for split, refs in data.items():
            for tau in (None, *TAUS):
                filename = f"{split}-forced.jsonl" if tau is None else f"{split}-t{tau:.2f}.jsonl"
                path = folder / filename
                rows = load(path, True, done["outputs"][filename])
                if [r["id"] for r in rows] != [r["id"] for r in refs]:
                    raise ValueError("Prediction IDs/order differ from references")
                for line, (row, ref) in enumerate(zip(rows, refs), 1):
                    if (row["variant"] != variant or row["tau"] != tau or row["evaluation_split"] != split
                            or digest({k: v for k, v in row.items() if k != "row_sha256"}) != row["row_sha256"]):
                        raise ValueError("Row identity/checksum mismatch")
                    occurrence = {"variant": variant, "arm": row["arm"], "seed": row["seed"], "tau": tau,
                                  "split": split, "mode": row["mode"], "automatic_outcome": row["outcome"],
                                  "source_file": str(path.relative_to(ROOT)), "source_line": line,
                                  "row_sha256": row["row_sha256"]}
                    unit = add_occurrence(units, row, ref, occurrence, row["arm"] in ARMS and tau in TAUS)
                    if variant == "original" and tau is None:
                        shared[split, row["id"]] = unit["unit_key"]
    original_logits = {}
    for seed in SEEDS:
        folder = artifacts / f"selector/test-evaluation/s{seed}"
        done = load(folder / "complete.json")
        if done["status"] != "complete" or done["evaluation_lock_sha256"] != lock_sha:
            raise ValueError("Incomplete or unrelated selector receipt")
        for split, refs in data.items():
            path = folder / f"{split}-scores.json"
            envelope = load(path, expected=done["files"][path.name])
            if (envelope["seed"] != seed or envelope["split"] != split or envelope["identity_sha256"] != done["identity_sha256"]
                    or digest({k: v for k, v in envelope.items() if k != "sha256"}) != envelope["sha256"]):
                raise ValueError("Selector envelope identity/checksum mismatch")
            if [r["id"] for r in envelope["rows"]] != [r["id"] for r in refs]:
                raise ValueError("Selector population mismatch")
            for index, (row, ref) in enumerate(zip(envelope["rows"], refs)):
                key = shared[split, row["id"]]
                unit = units[key]
                if unit_key(row) != key or grade(row["text"], ref["aliases"], row["truncated"]).outcome != row["outcome"]:
                    raise ValueError("Selector does not use the shared candidate")
                original = row["logits"]["original_calibrated"]
                if key in original_logits and original_logits[key] != original:
                    raise ValueError("Deterministic filter differs across critic seeds")
                original_logits[key] = original
                for tau in TAUS:
                    unit["selector_pairs"].append({"seed": seed, "tau": tau,
                        "original_emits": emits(row, original, tau),
                        "critic_emits": emits(row, row["logits"]["critic_calibrated"], tau),
                        "original_logit": original, "critic_logit": row["logits"]["critic_calibrated"],
                        "source_file": str(path.relative_to(ROOT)), "row_index_zero_based": index})
    for key in shared.values():
        if len(units[key]["selector_pairs"]) != 6:
            raise ValueError("Missing paired filter decisions")
    for path in (Path(__file__), ROOT / "tests/test_audit_v2.py", ROOT / "src/abstention/io.py", ROOT / "src/abstention/scoring.py"):
        sources[str(path.relative_to(ROOT))] = file_digest(path)
    verify_sources(sources)
    return units, sources


def frame_and_stratum(unit):
    if unit["selector_pairs"]:
        pattern = decision_pattern(unit["selector_pairs"])
        block = "shared_agreement" if pattern.startswith("agreement_") else "shared_disagreement"
        return (unit["split"], block), (pattern, unit["automatic_outcome"])
    if unit["rl_representative"] is not None:
        return (unit["split"], "rl"), (unit["rl_representative"]["arm"], unit["automatic_outcome"])
    return None


def sample_units(units):
    strata = defaultdict(list)
    for key, unit in units.items():
        result = frame_and_stratum(unit)
        if result is not None:
            block, stratum = result
            strata[(*block, *stratum)].append(key)
    for stratum, keys in strata.items():
        keys.sort(key=lambda key: digest([SEED, "unit", stratum, key]))
    counts = Counter()
    for split in SPLITS:
        for block, quota in POLICY["per_dataset"].items():
            relevant = sorted((s for s in strata if s[:2] == (split, block)), key=lambda s: digest([SEED, "stratum", s]))
            if sum(len(strata[s]) for s in relevant) < quota:
                raise ValueError(f"Insufficient units for {split}/{block}; fixed quota {quota}; no automatic retuning")
            n = 0
            while n < quota:
                for s in relevant:
                    if counts[s] < len(strata[s]) and n < quota:
                        counts[s] += 1; n += 1
    selected, summary = [], []
    for s in sorted(strata):
        N, n = len(strata[s]), counts[s]
        summary.append({"stratum": list(s), "population_units": N, "sample_units": n,
                        "nominal_inclusion_fraction": n/N})
        for key in strata[s][:n]:
            selected.append({**units[key], "sampling": {"stratum": list(s), "N_h": N, "n_h": n,
                                                        "nominal_inclusion_fraction": n/N, "design_weight": N/n}})
    selected.sort(key=lambda unit: digest([SEED, "blind_order", unit["unit_key"]]))
    if len(selected) != 150 or len({u["unit_key"] for u in selected}) != 150:
        raise AssertionError("Sample does not contain exactly150 unique units")
    return selected, {"strata": summary, "unique_units": len(selected),
                      "unique_questions": len({u["question_id"] for u in selected}),
                      "sample_by_block": dict(Counter("/".join(u["sampling"]["stratum"][:2]) for u in selected)),
                      "representative_rl_seeds": dict(Counter(str(u["rl_representative"]["seed"]) for u in selected if u["sampling"]["stratum"][1] == "rl"))}


def blind_rows(selected):
    return [{"audit_id": f"V2-{i:03d}", "question": u["question"], "response": u["response"],
             "reference_aliases": "\n".join(u["aliases"]), "truncated": "si" if u["truncated"] else "no",
             "human_outcome": "", "format_note": "", "evidence": "", "notes": ""}
            for i, u in enumerate(selected, 1)]


GUIDE = [
    ("Qué hacer", "En Auditoria, completa human_outcome: correct, error, abstain o unclear. Conserva las columnas A–E. Puedes ordenar filas; no cambies IDs ni textos."),
    ("correct / error", "Juzga el significado: correct si responde correctamente; error si es falsa, contradictoria o insuficiente. Acepta nombres equivalentes y paráfrasis. Las referencias orientan, pero no son exhaustivas ni infalibles."),
    ("abstain / unclear", "abstain si se abstiene sin dar una respuesta, incluso con lenguaje natural. unclear si no puedes resolver el caso o falta contexto. No adivines: explica la duda en notes."),
    ("Formato separado", "truncated es un dato observado, no tu veredicto. Un texto truncado todavía puede expresar una respuesta completa. format_note es OPCIONAL: ok, incomplete, empty, abstention_variant, other o unclear. Una respuesta vacía es error semántico."),
    ("Evidencia y tiempo", "Si consultas fuentes, pega URL o referencia en evidence y la razón en notes. No sustituyas silenciosamente la fecha/contexto de la pregunta por el presente. Si sigue ambiguo, usa unclear."),
    ("Revisión personal", "Revisa personalmente cada etiqueta. Completa Procedencia, declarando cualquier apoyo de IA y cómo lo usaste. El archivo comienza sin etiquetas ni método de anotación supuesto."),
    ("Ceguera y alcance", "No consultes claves privadas, decisiones automáticas ni el muestreo durante la revisión. Esta muestra diagnóstica no estima directamente una tasa poblacional de error."),
    ("Al terminar", "Guarda una copia XLSX con tus respuestas y devuélvela. No necesitas abrir archivos JSON. El importador comprobará IDs, textos, etiquetas y procedencia sin cambiar los resultados originales."),
]


def write_workbook(path, records):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    if path.exists():
        raise FileExistsError("Never overwrite a possibly annotated workbook")
    book = Workbook()
    guide = book.active; guide.title = "Instrucciones"
    audit = book.create_sheet("Auditoria"); provenance = book.create_sheet("Procedencia")
    book.properties.creator = "Auditoría de respuestas"
    book.properties.title = "Revisión ciega de respuestas"
    book.properties.subject = book.properties.description = book.properties.keywords = ""
    guide.append(("Guía breve", "Instrucciones"))
    for row in GUIDE:
        guide.append(row)
    audit.append(FIELDS)
    for record in records:
        if any(len(record[k]) > 32767 for k in FIELDS):
            raise ValueError("XLSX cell limit exceeded; refuse silent truncation")
        audit.append([record[k] for k in FIELDS])
    provenance.append(("Campo", "Respuesta"))
    for question in PROVENANCE:
        provenance.append((question, ""))
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                cell.data_type = "s"  # Literal dataset strings, including leading '='.
                cell.number_format = "@"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(size=11)
            sheet.row_dimensions[row[0].row].height = 80 if sheet == audit else 66
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="203864")
            cell.font = Font(color="FFFFFF", bold=True)
    for col, width in zip("ABCDEFGHI", (14, 55, 40, 52, 12, 20, 23, 42, 48)):
        audit.column_dimensions[col].width = width
    for row in audit.iter_rows(min_row=2, min_col=6, max_col=9):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for sheet in (guide, provenance):
        sheet.column_dimensions["A"].width = 42; sheet.column_dimensions["B"].width = 108
        sheet.freeze_panes = "B2"
    for row in provenance.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].fill = PatternFill("solid", fgColor="FFF2CC")
    audit.freeze_panes = "C2"; audit.auto_filter.ref = f"A1:I{len(records)+1}"
    for column, choices in (("F", LABELS), ("G", FORMATS)):
        dv = DataValidation(type="list", formula1='"'+','.join(choices)+'"', allow_blank=True)
        dv.showErrorMessage = True; dv.errorTitle = "Valor no válido"; dv.error = "Elige una opción de la lista."
        audit.add_data_validation(dv); dv.add(f"{column}2:{column}{len(records)+1}")
    dv = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
    dv.showErrorMessage = True; provenance.add_data_validation(dv); dv.add("B4"); dv.add("B6")
    book.save(path)
    verify_workbook_structure(load_workbook(path, data_only=False))


def verify_workbook_structure(book):
    if book.sheetnames != ["Instrucciones", "Auditoria", "Procedencia"]:
        raise ValueError("Unexpected, missing or extra workbook sheets")
    if book.custom_doc_props:
        raise ValueError("Unexpected custom metadata")
    for sheet in book:
        if sheet.sheet_state != "visible" or any(d.hidden for d in sheet.row_dimensions.values()) or any(d.hidden for d in sheet.column_dimensions.values()):
            raise ValueError("Hidden rows, columns or sheets are forbidden")
        if any(c.data_type == "f" or c.comment is not None for row in sheet for c in row):
            raise ValueError("Executable formulas or comments are forbidden")
    if tuple(c.value for c in book["Auditoria"][1]) != FIELDS or book["Auditoria"].max_column != len(FIELDS):
        raise ValueError("Audit headers/columns changed")


def verify_sources(sources):
    for relative, checksum in sources.items():
        if file_digest(ROOT / relative) != checksum:
            raise ValueError(f"Source changed: {relative}")


def csv_literal(value):
    # Excel-safe backup; the XLSX and private manifest retain exact unescaped text.
    return "'"+value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def export_audit(output=DEFAULT_OUT, loader=load_population):
    output = Path(output)
    if output.exists():
        raise FileExistsError("Audit destination exists; refusing to overwrite or silently resample")
    units, sources = loader()
    selected, sampling = sample_units(units)
    records = blind_rows(selected)
    verify_sources(sources)
    output.mkdir(parents=True, exist_ok=False)
    # No labels, reviewer identity or prior-audit annotation provenance are inferred.
    write_json(output / "private-key.json", {"policy_sha256": digest(POLICY),
               "units": {r["audit_id"]: u for r, u in zip(records, selected)}, "blind_records": records})
    write_json(output / "sampling.json", {"created_at": utcnow(), "policy": POLICY, **sampling, "source_sha256": sources})
    with (output / "blind.csv").open("x", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS); writer.writeheader()
        writer.writerows({k: csv_literal(v) for k, v in r.items()} for r in records)
    write_workbook(output / "auditoria-v2-150-respuestas.xlsx", records)
    verify_sources(sources)
    receipt = {"status": "blank_audit_ready", "at": utcnow(), "human_labels_completed": False,
               "policy_sha256": digest(POLICY), "units": 150, "unique_questions": sampling["unique_questions"],
               "workbook": "auditoria-v2-150-respuestas.xlsx", "csv_formula_prefix_escaping": True,
               "files": {p.name: file_digest(p) for p in sorted(output.iterdir()) if p.is_file()},
               "source_sha256": sources}
    write_json(output / "export-receipt.json", receipt)
    return receipt


def validate_submission(path, audit_dir=DEFAULT_OUT):
    from openpyxl import load_workbook
    audit_dir, path = Path(audit_dir), Path(path)
    receipt = read_json(audit_dir / "export-receipt.json")
    if receipt["status"] != "blank_audit_ready" or receipt["policy_sha256"] != digest(POLICY):
        raise ValueError("Wrong audit protocol")
    for filename, checksum in receipt["files"].items():
        if file_digest(audit_dir / filename) != checksum:
            raise ValueError("Original audit export changed; return a separate annotated copy")
    mapping = read_json(audit_dir / "private-key.json")
    if mapping["policy_sha256"] != digest(POLICY):
        raise ValueError("Private mapping belongs to another protocol")
    original = {r["audit_id"]: r for r in mapping["blind_records"]}
    if len(original) != 150 or set(original) != set(mapping["units"]):
        raise ValueError("Invalid original sample mapping")
    before = file_digest(path)
    book = load_workbook(path, data_only=False, keep_links=False)
    verify_workbook_structure(book)
    sheet = book["Auditoria"]
    if sheet.max_row != 151:
        raise ValueError("Expected exactly150 annotation rows")
    submitted, seen = [], set()
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(FIELDS, ("" if v is None else v for v in cells)))
        if any(not isinstance(v, str) for v in row.values()):
            raise ValueError("Annotation cells must be text")
        audit_id = row["audit_id"]
        if audit_id in seen or audit_id not in original:
            raise ValueError("Duplicate or unknown audit ID")
        seen.add(audit_id)
        if any(row[k] != original[audit_id][k] for k in FIELDS[:5]):
            raise ValueError("Question/response/reference/truncation fields were modified")
        if row["human_outcome"] not in LABELS:
            raise ValueError(f"Missing or invalid semantic label: {audit_id}")
        if row["format_note"] and row["format_note"] not in FORMATS:
            raise ValueError("Invalid optional format note")
        if row["human_outcome"] == "unclear" and not (row["notes"].strip() or row["evidence"].strip()):
            raise ValueError("Explain an unclear label in notes or evidence")
        submitted.append(row)
    if seen != set(original):
        raise ValueError("Missing audit IDs")
    provenance_sheet = book["Procedencia"]
    if provenance_sheet.max_row != len(PROVENANCE)+1 or provenance_sheet.max_column != 2:
        raise ValueError("Annotation provenance structure changed")
    provenance = {}
    for index, question in enumerate(PROVENANCE, 2):
        if provenance_sheet.cell(index, 1).value != question:
            raise ValueError("Provenance question changed")
        value = provenance_sheet.cell(index, 2).value
        if value is not None and not isinstance(value, str):
            raise ValueError("Provenance must be text")
        provenance[question] = (value or "").strip()
    if not provenance[PROVENANCE[0]]:
        raise ValueError("Annotator identifier is required")
    try:
        date.fromisoformat(provenance[PROVENANCE[1]])
    except ValueError as exc:
        raise ValueError("Review date must be YYYY-MM-DD") from exc
    for key in (PROVENANCE[2], PROVENANCE[4]):
        if provenance[key] not in {"si", "sí", "no"}:
            raise ValueError("Declare AI assistance and personal review explicitly")
    if provenance[PROVENANCE[2]] in {"si", "sí"} and not provenance[PROVENANCE[3]]:
        raise ValueError("Describe the AI tool and how it was used")
    if file_digest(path) != before:
        raise ValueError("Submission changed while being read")
    submitted.sort(key=lambda r: r["audit_id"])
    return submitted, provenance, mapping, receipt, before


def import_audit(path, audit_dir=DEFAULT_OUT):
    path, audit_dir = Path(path), Path(audit_dir)
    rows, provenance, mapping, export, checksum = validate_submission(path, audit_dir)
    destination = audit_dir / "submissions" / checksum[:16]
    if destination.exists():
        raise FileExistsError("This submission was already imported; no overwrite")
    destination.mkdir(parents=True, exist_ok=False)
    shutil.copy2(path, destination / "submitted.xlsx")
    joined = [{**r, "unit_key": mapping["units"][r["audit_id"]]["unit_key"],
               "automatic_outcome": mapping["units"][r["audit_id"]]["automatic_outcome"],
               "sampling": mapping["units"][r["audit_id"]]["sampling"]} for r in rows]
    write_json(destination / "labels.json", {"annotations": joined, "provenance": provenance})
    confusion = Counter((r["automatic_outcome"], r["human_outcome"]) for r in joined)
    result = {"status": "labels_imported_pending_adjudication", "at": utcnow(), "n": len(rows),
              "unclear": sum(r["human_outcome"] == "unclear" for r in rows),
              "personal_review_all_declared": provenance[PROVENANCE[4]] in {"si", "sí"},
              "ai_assistance_declared": provenance[PROVENANCE[2]] in {"si", "sí"},
              "confusion_diagnostic_only": [{"automatic": a, "semantic_human": h, "n": n} for (a, h), n in sorted(confusion.items())],
              "submission_sha256": checksum, "export_receipt_sha256": file_digest(audit_dir / "export-receipt.json"),
              "labels_sha256": file_digest(destination / "labels.json"),
              "frozen_endpoints_modified": False,
              "interpretation": POLICY["limits"]}
    if file_digest(destination / "submitted.xlsx") != checksum or file_digest(path) != checksum:
        raise ValueError("Submission archive integrity failure")
    write_json(destination / "import-receipt.json", result)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("export", "validate", "import"))
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--workbook", type=Path)
    args = parser.parse_args()
    if args.command == "export":
        result = export_audit(args.output)
    else:
        if args.workbook is None:
            parser.error("--workbook is required")
        if args.command == "validate":
            rows, provenance, _, _, checksum = validate_submission(args.workbook, args.output)
            result = {"status": "valid", "n": len(rows), "submission_sha256": checksum,
                      "unclear": sum(r["human_outcome"] == "unclear" for r in rows)}
        else:
            result = import_audit(args.workbook, args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
